"""Exports (Excel, PDF, e-mails Outlook) et contrôles qualité des fichiers importés."""
import html as html_lib
import re
from email.message import EmailMessage
from io import BytesIO

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from metier import (COLS_ATTRIBUTION, COLS_LISTES, COLS_VOYAGE, VILLES_EXCLUES_ANALYSE,
                    libelle_besoin, regle_zone7, trier_par_zone)

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# =====================================================
# CONTRÔLES QUALITÉ (poids 0/1000, volumes nuls)
# =====================================================
POIDS_SUSPECTS = [0, 1000]
CLIENTS_INTERNES = {"QUIVEM", "SOPALTHYNA", "SOPAL", "SANA"}
SITES_VOLUME = {"SA1", "AP1", "SO1", "TC1", "SO2"}
CLIENTS_EXCLUS_VOLUME = {
    "AMECAP", "SANA", "SOPAL", "SOPALGAZ", "SOPALSERV", "SOPALTEC", "SOPALALG", "AQUAPLAST",
    "WINOX", "QUIVEM", "SANISTONE", "SOPAMAR", "SOPALAFR", "SOPALINTER", "PERSOGSO", "SOPALTHYNA",
}

EMAIL_A = "Malek.ABDELLAOUI@sopal.com"
EMAIL_CC = ["Taha.BOUAZIZ@sopal.com", "Nadhem.KESKES@sopal.com", "Bechir.KHEMAKHEM@sopal.com"]
SIGNATURE = "Zaineb KCHAOU"


def detecter_poids_suspects(df_liv_raw, df_art_raw):
    """Lignes dont 'Poids de l'US' vaut 0 ou 1000, hors BL des clients internes."""
    df = df_liv_raw.copy()
    df["Poids de l'US"] = pd.to_numeric(df["Poids de l'US"], errors="coerce")
    if "Client commande" in df.columns:
        df = df[~df["Client commande"].astype(str).str.strip().str.upper().isin(CLIENTS_INTERNES)]
    suspects = df[df["Poids de l'US"].isin(POIDS_SUSPECTS)][["No livraison", "Article", "Poids de l'US"]].copy()

    col_design = next((c for c in df_art_raw.columns if "signation" in str(c)), None)
    if col_design is not None and "Article" in df_art_raw.columns:
        art = df_art_raw[["Article", col_design]].drop_duplicates(subset="Article")
        art = art.rename(columns={col_design: "Désignation"})
        suspects["Article"] = suspects["Article"].astype(str).str.strip()
        art["Article"] = art["Article"].astype(str).str.strip()
        suspects = suspects.merge(art, on="Article", how="left")
    else:
        suspects["Désignation"] = ""
    suspects = suspects.rename(columns={"Poids de l'US": "Poids"})
    return suspects[["No livraison", "Article", "Désignation", "Poids"]].drop_duplicates().reset_index(drop=True)


def detecter_volumes_nuls(df_liv_raw, df_art_raw):
    """Articles (sites autorisés, hors SDC et clients internes) au volume nul ou absent."""
    df = df_liv_raw.copy()
    df = df[df["Site vente"].astype(str).str.strip().str.upper().isin(SITES_VOLUME)]
    df = df[df["Type livraison"].astype(str).str.strip().str.upper() != "SDC"]
    df = df[~df["Client commande"].astype(str).str.strip().str.upper().isin(CLIENTS_EXCLUS_VOLUME)]
    df["Article"] = df["Article"].astype(str).str.strip()
    articles = df.groupby("Article")["No livraison"].nunique().reset_index(name="Nb BL")

    art = df_art_raw.copy()
    col_vol = next((c for c in art.columns if "olume" in str(c)), None)
    col_des = next((c for c in art.columns if "signation" in str(c)), None)
    if col_vol is None:
        raise ValueError(f"Colonne volume introuvable. Colonnes : {list(art.columns)}")
    art["Article"] = art["Article"].astype(str).str.strip()
    art["_vol"] = pd.to_numeric(art[col_vol], errors="coerce")
    agg = {"_vol": "max"}
    if col_des is not None:
        agg[col_des] = "first"
    art = art.groupby("Article", as_index=False).agg(agg)

    res = articles.merge(art, on="Article", how="left").rename(columns={"_vol": "Volume"})
    if col_des is not None:
        res = res.rename(columns={col_des: "Désignation"})
    else:
        res["Désignation"] = ""
    res = res[res["Volume"].isna() | (res["Volume"] == 0)]
    return res[["Article", "Désignation", "Volume", "Nb BL"]].sort_values("Article").reset_index(drop=True)


# =====================================================
# MOTEUR D'EXPORT EXCEL (mise en forme unique)
# =====================================================
_COTE = Side(style="thin", color="B0C4DE")
_BORDURE = Border(left=_COTE, right=_COTE, top=_COTE, bottom=_COTE)
_CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
_FOND_ENTETE = PatternFill(fill_type="solid", fgColor="0369A1")

LARGEURS = {
    "No livraison": 18, "Client": 18, "Ville": 18, "Zone": 14, "Représentant": 18, "Article": 24,
    "Désignation": 45, "Poids total": 16, "Volume total": 16, "Nombre de BLs": 16,
    "Besoin estafette réel": 30, "Véhicule N°": 15, "Poids total chargé": 20,
    "Volume total chargé": 21, "Client(s) inclus": 32, "Représentant(s) inclus": 30,
    "BL inclus": 35, "Taux occupation poids (%)": 24, "Taux occupation volume (%)": 26,
    "Taux d'occupation (%)": 23, "Location_camion": 18, "Location_proposee": 20,
    "Code Véhicule": 18, "Véhicule attribué": 22, "Chauffeur attribué": 25,
    "Matricule chauffeur": 20, "Date Voyage Format": 18, "Code voyage": 26,
}
FORMATS = {
    "Poids total": "0.000", "Volume total": "0.000",
    "Poids total chargé": "0.000", "Volume total chargé": "0.000",
    "Taux occupation poids (%)": '0.000"%"', "Taux occupation volume (%)": '0.000"%"',
    "Taux d'occupation (%)": '0.000"%"',
}


def _mettre_en_forme(ws, df, hauteur_fixe=None):
    for ligne in ws.iter_rows():
        for c in ligne:
            c.border = _BORDURE
            c.alignment = _CENTRE
            c.font = Font(name="Arial", size=10)
    for c in ws[1]:
        c.fill = _FOND_ENTETE
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    ws.row_dimensions[1].height = 32

    for i, nom in enumerate(df.columns, start=1):
        lettre = get_column_letter(i)
        if nom in LARGEURS:
            ws.column_dimensions[lettre].width = LARGEURS[nom]
        else:
            longueur = max((max(len(l) for l in str(v).split("\n")) for v in df[nom] if v is not None), default=10)
            ws.column_dimensions[lettre].width = min(max(longueur, len(str(nom))) + 2, 40)
        if nom in FORMATS:
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=i).number_format = FORMATS[nom]

    for r in range(2, ws.max_row + 1):
        if hauteur_fixe:
            ws.row_dimensions[r].height = hauteur_fixe
            continue
        nb = max((str(c.value).count("\n") + 1 for c in ws[r] if c.value is not None), default=1)
        ws.row_dimensions[r].height = min(max(22, nb * 18), 120)

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def excel(feuilles):
    """feuilles : liste de (nom_feuille, DataFrame[, hauteur_fixe])."""
    sortie = BytesIO()
    with pd.ExcelWriter(sortie, engine="openpyxl") as writer:
        for f in feuilles:
            nom, df = f[0][:31], f[1]
            df.to_excel(writer, index=False, sheet_name=nom)
            _mettre_en_forme(writer.sheets[nom], df, f[2] if len(f) > 2 else None)
    return sortie.getvalue()


def _une_par_ligne(df, cols, sep=r"[;,]"):
    df = df.copy()
    for c in cols:
        if c in df.columns:
            df[c] = df[c].map(lambda v: "\n".join(
                e.strip() for e in re.split(sep, "" if pd.isna(v) else str(v)) if e.strip() and e.strip() != "nan"))
    return df


def _arrondir(df, cols, n=3):
    df = df.copy()
    for c in cols:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").round(n)
    return df


# =====================================================
# PRÉPARATION DES DONNÉES À EXPORTER
# =====================================================
def prep_livraisons(df, sans_zone=True):
    d = df.copy()
    if sans_zone:
        d = d.drop(columns=["Zone"], errors="ignore")
        d = d[~d["Ville"].isin(VILLES_EXCLUES_ANALYSE)]
    d = _une_par_ligne(d, ["Article"], r",")
    d = _une_par_ligne(d, [c for c in d.columns if "signation" in str(c)], r"<br\s*/?>|\|\|")
    return _arrondir(d, ["Poids total", "Volume total"])


def prep_besoin(df, par_zone=False):
    d = df.rename(columns={"Nombre livraisons": "Nombre de BLs"}).copy()
    if not par_zone:
        d = d[~d["Ville"].isin(VILLES_EXCLUES_ANALYSE)]
    d = _arrondir(d, ["Poids total", "Volume total"])
    if "Besoin estafette réel" in d.columns:
        def besoin(r):
            v = r["Besoin estafette réel"]
            if regle_zone7(r):
                return libelle_besoin(v, True, sep="\n" if par_zone else " → ")
            v = pd.to_numeric(v, errors="coerce")
            return None if pd.isna(v) else round(float(v), 1)
        d["Besoin estafette réel"] = d.apply(besoin, axis=1)
    return d


def prep_voyages(df, colonnes=None, extra=()):
    colonnes = [c for c in (colonnes or COLS_VOYAGE) + list(extra) if c in df.columns]
    d = trier_par_zone(df[colonnes].copy())
    d = _une_par_ligne(d, COLS_LISTES)
    return _arrondir(d, ["Poids total chargé", "Volume total chargé", "Taux occupation poids (%)",
                         "Taux occupation volume (%)", "Taux d'occupation (%)"])


@st.cache_data(show_spinner=False)
def excel_livraisons(df, sans_zone=True, nom="Livraisons"):
    return excel([(nom, prep_livraisons(df, sans_zone))])


@st.cache_data(show_spinner=False)
def excel_besoin(df, par_zone=False, nom="Besoin"):
    return excel([(nom, prep_besoin(df, par_zone), None if par_zone else 30)])


@st.cache_data(show_spinner=False)
def excel_voyages(df, nom="Voyages", colonnes=None, extra=()):
    return excel([(nom, prep_voyages(df, colonnes, extra))])


@st.cache_data(show_spinner=False)
def excel_analyse_complete(df_grouped, df_city, df_grouped_zone, df_zone):
    return excel([
        ("Livraisons Client Ville", prep_livraisons(df_grouped, True)),
        ("Besoin par Ville", prep_besoin(df_city, False), 30),
        ("Livraisons Client Zone", prep_livraisons(df_grouped_zone, False)),
        ("Besoin par Zone", prep_besoin(df_zone, True)),
    ])


COLS_ATTRIBUES = ["Zone", "Véhicule N°", "Poids total chargé", "Volume total chargé", "Client(s) inclus",
                  "Représentant(s) inclus", "BL inclus", "Taux d'occupation (%)"] + COLS_ATTRIBUTION
COLS_CODES = COLS_VOYAGE + ["Taux d'occupation (%)", "Type_Camion"] + COLS_ATTRIBUTION + \
    ["Date Voyage Format", "Code voyage"]


# =====================================================
# LISTES À CORRIGER (poids / volumes) + E-MAILS OUTLOOK
# =====================================================
def _excel_correction(df, nom, col_rouge, largeurs):
    sortie = BytesIO()
    with pd.ExcelWriter(sortie, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=nom)
        ws = writer.sheets[nom]
        _mettre_en_forme(ws, df)
        for i, c in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = largeurs.get(c, 18)
        j = list(df.columns).index(col_rouge) + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=j)
            cell.number_format = "0.000"
            cell.fill = PatternFill(fill_type="solid", fgColor="FDE2E2")
            cell.font = Font(name="Arial", size=10, bold=True, color="B91C1C")
    return sortie.getvalue()


def excel_poids_suspects(df):
    d = df.copy()
    d["Poids"] = pd.to_numeric(d["Poids"], errors="coerce")
    return _excel_correction(d, "Poids à corriger", "Poids",
                             {"No livraison": 22, "Article": 18, "Désignation": 50, "Poids": 14})


def excel_volumes_nuls(df):
    return _excel_correction(df, "Volumes à corriger", "Volume",
                             {"Article": 18, "Désignation": 55, "Volume": 14, "Nb BL": 10})


def _email_html(intro, entetes, largeurs, lignes):
    cellule = "border:1px solid #B0C4DE;padding:7px 10px;font-family:Calibri,Arial,sans-serif;font-size:10.5pt;"
    th = "".join(
        f'<td width="{largeurs[h]}" align="center" bgcolor="#0369A1" style="background-color:#0369A1;{cellule}'
        f'text-align:center;"><font color="#FFFFFF"><b>{h}</b></font></td>\n' for h in entetes)
    corps = ""
    for i, valeurs in enumerate(lignes):
        fond = "#FFFFFF" if i % 2 == 0 else "#F0F7FD"
        tds = ""
        for h, (v, align, rouge) in zip(entetes, valeurs):
            contenu = f'<font color="#B91C1C"><b>{v}</b></font>' if rouge else v
            tds += (f'<td width="{largeurs[h]}" align="{align}" bgcolor="{fond}" '
                    f'style="{cellule}text-align:{align};">{contenu}</td>\n')
        corps += f'<tr bgcolor="{fond}">\n{tds}</tr>\n'
    largeur = sum(largeurs.values())
    return f"""<html><head><meta http-equiv="Content-Type" content="text/html; charset=utf-8"></head>
<body style="font-family:Calibri,Arial,sans-serif;font-size:11pt;color:#000000;">
<p>Bonjour,</p>
<p>{intro}</p>
<table width="{largeur}" cellpadding="0" cellspacing="0" border="0" style="border-collapse:collapse;width:{largeur}px;">
<tr bgcolor="#0369A1">
{th}</tr>
{corps}</table>
<p>Le fichier Excel de cette liste est joint à ce message.<br>
Merci de me confirmer une fois les corrections effectuées afin que je relance le planning.</p>
<p>Cordialement,<br><b>{SIGNATURE}</b><br></p>
</body></html>"""


def _eml(sujet, texte, corps_html, piece, nom_piece):
    msg = EmailMessage()
    msg["To"] = EMAIL_A
    msg["Cc"] = ", ".join(EMAIL_CC)
    msg["Subject"] = sujet
    msg["X-Unsent"] = "1"  # s'ouvre comme brouillon dans Outlook
    msg.set_content(texte, cte="base64")
    msg.add_alternative(corps_html, subtype="html", cte="base64")
    msg.add_attachment(piece, maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=nom_piece)
    return bytes(msg)


def eml_poids_suspects(df):
    n, nb_bl = len(df), df["No livraison"].nunique()
    lignes = []
    for _, r in df.iterrows():
        p = pd.to_numeric(r["Poids"], errors="coerce")
        lignes.append([(html_lib.escape(str(r["No livraison"])), "center", False),
                       (html_lib.escape(str(r["Article"])), "center", False),
                       (html_lib.escape(str(r["Désignation"])), "left", False),
                       ("" if pd.isna(p) else f"{p:,.3f}".replace(",", " "), "center", True)])
    intro = (f"Lors du contrôle des livraisons, <b>{n} article(s)</b> répartis sur <b>{nb_bl} BL</b> ont été "
             f"détectés avec un <b>poids égal à 0 ou 1000</b>.<br>Merci de vérifier et de corriger le poids de "
             f"ces articles dans la fiche article :")
    corps = _email_html(intro, ["No livraison", "Article", "Désignation", "Poids"],
                        {"No livraison": 150, "Article": 120, "Désignation": 380, "Poids": 90}, lignes)
    texte = (f"Bonjour,\n\n{n} article(s) ont un poids égal à 0 ou 1000. Merci de corriger leur poids "
             f"(voir pièce jointe).\n\nCordialement,\n{SIGNATURE}")
    return _eml(f"Articles avec poids à corriger (0 ou 1000) - {n} article(s)", texte, corps,
                excel_poids_suspects(df), "articles_poids_a_corriger.xlsx")


def eml_volumes_nuls(df):
    n = len(df)
    lignes = []
    for _, r in df.iterrows():
        v = pd.to_numeric(r["Volume"], errors="coerce")
        lignes.append([(html_lib.escape(str(r["Article"])), "center", False),
                       (html_lib.escape(str(r["Désignation"])), "left", False),
                       ("Non renseigné" if pd.isna(v) else f"{v:.3f}", "center", True),
                       (str(int(r["Nb BL"])), "center", False)])
    intro = (f"Lors du contrôle des livraisons, <b>{n} article(s)</b> ont été détectés avec un <b>volume nul ou "
             f"non renseigné</b>.<br>Merci de vérifier et de corriger le volume de ces articles dans la fiche article :")
    corps = _email_html(intro, ["Article", "Désignation", "Volume", "Nb BL"],
                        {"Article": 130, "Désignation": 420, "Volume": 100, "Nb BL": 80}, lignes)
    texte = (f"Bonjour,\n\n{n} article(s) ont un volume nul ou non renseigné. Merci de corriger leur volume "
             f"(voir pièce jointe).\n\nCordialement,\n{SIGNATURE}")
    return _eml(f"Articles avec volume nul à corriger - {n} article(s)", texte, corps,
                excel_volumes_nuls(df), "articles_volume_nul.xlsx")


# =====================================================
# EXPORTS PDF
# =====================================================
def pdf_attribues(df, title="Voyages Attribués"):
    from fpdf import FPDF
    pdf = FPDF(orientation='L')
    pdf.add_page()

    pdf.set_left_margin(5)
    pdf.set_right_margin(5)
    pdf.set_top_margin(10)

    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 8, title, ln=True, align="C")
    pdf.ln(3)

    df_pdf = df.copy()

    numeric_columns = {
        'Poids total chargé': ('kg', 3),
        'Volume total chargé': ('m³', 3),
        'Taux d\'occupation (%)': ('%', 2)
    }

    for col, (unit, decimals) in numeric_columns.items():
        if col in df_pdf.columns:
            df_pdf[col] = df_pdf[col].apply(
                lambda x: f"{float(x):.{decimals}f} {unit}" if x and str(x).strip() and str(x).strip() != 'nan' else ""
            )

    col_config = {
        'Zone': {'width': 30, 'header': 'Zone'},               # élargie (15 -> 30)
        'Véhicule N°': {'width': 15, 'header': 'Véhicule'},    # réduite (18 -> 15)
        'Poids total chargé': {'width': 22, 'header': 'Poids (kg)'},
        'Volume total chargé': {'width': 22, 'header': 'Volume (m³)'},
        'Client(s) inclus': {'width': 30, 'header': 'Clients'},
        'Représentant(s) inclus': {'width': 30, 'header': 'Représentants'},
        'BL inclus': {'width': 35, 'header': 'BL associés'},
        'Taux d\'occupation (%)': {'width': 18, 'header': 'Taux %'},
        'Véhicule attribué': {'width': 25, 'header': 'Véhicule Attribué'},
        'Chauffeur attribué': {'width': 25, 'header': 'Chauffeur'},
        'Matricule chauffeur': {'width': 20, 'header': 'Matricule'}
    }

    colonnes_existantes = [col for col in df_pdf.columns if col in col_config]
    widths = [col_config[col]['width'] for col in colonnes_existantes]
    headers = [col_config[col]['header'] for col in colonnes_existantes]

    # ✅ CENTRAGE : position de départ calculée selon la largeur du tableau
    total_width = sum(widths)
    start_x = (pdf.w - total_width) / 2

    def dessiner_entetes():
        pdf.set_x(start_x)
        pdf.set_font("Arial", 'B', 8)
        for i, header in enumerate(headers):
            pdf.cell(widths[i], 6, header, border=1, align='C')
        pdf.ln()
        pdf.set_font("Arial", '', 7)

    dessiner_entetes()

    list_columns = ['Client(s) inclus', 'Représentant(s) inclus', 'BL inclus']

    for _, row in df_pdf.iterrows():
        if pdf.get_y() > 180:
            pdf.add_page()
            dessiner_entetes()

        max_lines = 1
        list_contents = {}

        for col in list_columns:
            if col in colonnes_existantes:
                content = str(row[col]) if pd.notna(row[col]) and str(row[col]) != 'nan' else ""
                elements = [e.strip() for e in content.replace(';', ',').split(',') if e.strip()]
                list_contents[col] = elements
                max_lines = max(max_lines, len(elements))

        for line_idx in range(max_lines):
            if pdf.get_y() > 190:
                pdf.add_page()
                dessiner_entetes()

            pdf.set_x(start_x)

            for i, col in enumerate(colonnes_existantes):
                if col in list_columns:
                    elements = list_contents.get(col, [])
                    content = elements[line_idx] if line_idx < len(elements) else ""
                else:
                    if line_idx == 0:
                        content = str(row[col]) if pd.notna(row[col]) and str(row[col]) != 'nan' else ""
                    else:
                        content = ""

                border = 'LR'
                if line_idx == 0:
                    border += 'T'
                if line_idx == max_lines - 1:
                    border += 'B'

                pdf.cell(widths[i], 5, content, border=border, align='C')

            pdf.ln()

    return pdf.output(dest='S').encode('latin-1')


def pdf_codes(df):
    from fpdf import FPDF
    import pandas as pd
    import math

    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_margins(10, 10, 10)
    pdf.set_auto_page_break(auto=True, margin=10)
    pdf.add_page()

    # =====================================================
    # STYLE GÉNÉRAL
    # =====================================================
    bleu = (31, 78, 121)
    bleu_clair = (221, 235, 247)
    gris_clair = (245, 247, 250)
    gris_bordure = (200, 208, 218)
    blanc = (255, 255, 255)
    noir = (40, 40, 40)

    # =====================================================
    # TITRE
    # =====================================================
    pdf.set_text_color(*bleu)
    pdf.set_font("Arial", "B", 15)
    pdf.cell(
        0, 9,
        "LISTE DES VOYAGES AVEC CODES",
        ln=True,
        align="C"
    )

    pdf.set_draw_color(*bleu)
    pdf.set_line_width(0.6)
    pdf.line(10, pdf.get_y() + 1, 287, pdf.get_y() + 1)
    pdf.ln(5)

    # =====================================================
    # COLONNES DANS L'ORDRE DEMANDÉ
    # =====================================================
    headers = [
        "Zone",
        "Véhicule N°",
        "Poids total chargé",
        "Volume total chargé",
        "Client(s) inclus",
        "Représentant(s) inclus",
        "BL inclus",
        "Taux d'occupation (%)",
        "Code Véhicule",
        "Véhicule attribué",
        "Chauffeur attribué",
        "Matricule chauffeur",
        "Date Voyage Format",
        "Code voyage"
    ]

    # Total : 277 mm, adapté à A4 paysage avec marges de 10 mm
    widths = [
        13, 15, 18, 18, 28, 21, 40,
        18, 17, 18, 20, 19, 17, 15
    ]

    # =====================================================
    # PRÉPARATION DES DONNÉES
    # =====================================================
    def get_value(row, name):
        value = row.get(name, "")

        if value is None:
            return ""

        try:
            if pd.isna(value):
                return ""
        except (TypeError, ValueError):
            pass

        if isinstance(value, float):
            return f"{value:.1f}" if "Taux" in name else f"{value:.3f}"
        return str(value).strip()

    def prepare_value(value, column_index):
        value = str(value)

        # Listes : un élément par ligne
        if column_index in (4, 5, 6):
            value = value.replace("\r\n", "\n").replace("\r", "\n")
            value = value.replace(";", "\n")
            value = value.replace(", ", "\n")
            value = value.replace(",", "\n")

        return value.strip()

    def wrap_text(text, width):
        """Calcule les lignes réellement nécessaires dans la cellule."""
        result = []

        for paragraph in str(text).split("\n"):
            if not paragraph:
                result.append("")
                continue

            current = ""

            # Découpe aussi les identifiants exceptionnellement longs
            for word in paragraph.split():
                if pdf.get_string_width(word) > width:
                    if current:
                        result.append(current)
                        current = ""

                    part = ""
                    for char in word:
                        if pdf.get_string_width(part + char) > width and part:
                            result.append(part)
                            part = char
                        else:
                            part += char

                    current = part
                    continue

                candidate = word if not current else current + " " + word

                if pdf.get_string_width(candidate) <= width:
                    current = candidate
                else:
                    if current:
                        result.append(current)
                    current = word

            if current:
                result.append(current)

        return result or [""]

    # =====================================================
    # EN-TÊTE RÉPÉTÉ SUR CHAQUE PAGE
    # =====================================================
    def draw_header():
        pdf.set_font("Arial", "B", 6.5)
        pdf.set_text_color(*blanc)
        pdf.set_fill_color(*bleu)
        pdf.set_draw_color(*gris_bordure)
        pdf.set_line_width(0.2)

        x0 = 10
        y0 = pdf.get_y()
        header_height = 13

        for i, header in enumerate(headers):
            x = x0 + sum(widths[:i])

            pdf.set_xy(x, y0)
            pdf.rect(
                x, y0, widths[i], header_height,
                style="DF"
            )

            lines = wrap_text(header, widths[i] - 2)
            line_height = 3

            text_height = len(lines) * line_height
            text_y = y0 + max(0, (header_height - text_height) / 2)

            pdf.set_xy(x + 1, text_y)

            for line in lines:
                pdf.cell(
                    widths[i] - 2,
                    line_height,
                    line,
                    border=0,
                    align="C"
                )
                pdf.set_xy(pdf.get_x(), pdf.get_y())

                # Repositionnement vertical explicite
                pdf.set_xy(x + 1, pdf.get_y() + line_height)

        pdf.set_xy(10, y0 + header_height)

    # =====================================================
    # LIGNES DU TABLEAU
    # =====================================================
    def draw_row(values, row_number):
        font_size = 6
        line_height = 3.2
        padding = 1.2

        pdf.set_font("Arial", "", font_size)

        wrapped = [
            wrap_text(value, widths[i] - 2 * padding)
            for i, value in enumerate(values)
        ]

        max_lines = max(len(lines) for lines in wrapped)
        row_height = max(6, max_lines * line_height + 2 * padding)

        # Saut de page si nécessaire
        if pdf.get_y() + row_height > 195:
            pdf.add_page()
            draw_header()
            pdf.set_font("Arial", "", font_size)

        x0 = 10
        y0 = pdf.get_y()

        # Couleurs alternées
        fill = blanc if row_number % 2 == 0 else gris_clair

        pdf.set_draw_color(*gris_bordure)
        pdf.set_line_width(0.2)

        for i, lines in enumerate(wrapped):
            x = x0 + sum(widths[:i])

            # Toutes les cellules ont la même hauteur
            pdf.set_fill_color(*fill)
            pdf.rect(
                x, y0, widths[i], row_height,
                style="DF"
            )

            # Valeurs
            pdf.set_text_color(*noir)
            pdf.set_font("Arial", "", font_size)

            content_height = len(lines) * line_height
            text_y = y0 + max(
                padding,
                (row_height - content_height) / 2
            )

            pdf.set_xy(x + padding, text_y)

            for line in lines:
                pdf.cell(
                    widths[i] - 2 * padding,
                    line_height,
                    line,
                    border=0,
                    align="L" if i in (4, 5, 6) else "C"
                )
                pdf.set_xy(
                    x + padding,
                    pdf.get_y() + line_height
                )

        pdf.set_xy(x0, y0 + row_height)

    # =====================================================
    # GÉNÉRATION DU TABLEAU UNIQUE
    # =====================================================
    draw_header()

    for index, (_, row) in enumerate(df.iterrows()):
        values = [
            get_value(row, "Zone"),
            get_value(row, "Véhicule N°"),
            get_value(row, "Poids total chargé"),
            get_value(row, "Volume total chargé"),
            get_value(row, "Client(s) inclus"),
            get_value(row, "Représentant(s) inclus"),
            get_value(row, "BL inclus"),
            get_value(row, "Taux d'occupation (%)"),
            get_value(row, "Code Véhicule"),
            get_value(row, "Véhicule attribué"),
            get_value(row, "Chauffeur attribué"),
            get_value(row, "Matricule chauffeur"),
            get_value(row, "Date Voyage Format"),
            get_value(row, "Code voyage")
        ]

        values = [
            prepare_value(value, i)
            for i, value in enumerate(values)
        ]

        draw_row(values, index)

    # =====================================================
    # EXPORT
    # =====================================================
    return pdf.output(dest="S").encode("latin-1")
